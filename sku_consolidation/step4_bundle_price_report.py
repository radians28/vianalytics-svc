"""Step 4b - Excel report (Bahasa Indonesia) of bundle prices vs. component prices.

Input  : raw/sku_bundle/bundling_master.csv, bundling_price.csv   (same loaders as step 4)
Output : sku_master/laporan_harga_bundling.xlsx
  Sheet "Ringkasan"            counts per status / possible mismatch
  Sheet "Perlu Dicek"          only bundles with a possible mismatch
  Sheet "Semua Bundle"         every bundle
  Sheet "Keterangan"           legend of the columns and rules
Each report sheet has three sections:
  1. INFO BUNDLE       kode SKU, nama produk, harga bundle
  2. KOMPONEN          one row per component: kode SKU, nama produk, harga dasar, qty, total harga
  3. ANALISIS          possible mismatch:
       (1) Beli A gratis B   bundle price ~ price of ONE component with qty 1 (the other components come free)
       (2) Selisih terlalu besar   |bundle - sum of components| / lower of the two > threshold (10%);
                              not marked when the bundle already meets (1)
       (3) Kemungkinan salah harga bundle price ~ sum of the components with ONE component (qty > 1) counted
                              as a single unit (e.g. 37/390109905M = 1 unit of 37/020100501A while the bundle holds 6)
     (1) and (3) are only evaluated when the price gap is above the threshold.
"""
import argparse
import os

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from common import bundle
from common.config import SKU_MASTER
from step4_bundle_price_check import build_check

STATUS_ID = {"ok": "Sesuai", "pass": "Lolos (dikecualikan)", "price_gap": "Perlu dicek",
             "bundle_price_missing": "Harga bundle kosong", "component_price_missing": "Harga komponen kosong",
             "price_zero": "Harga nol tidak konsisten"}
NAVY, TEAL, ORANGE = "1F3864", "1F6E6E", "9C5700"
FILL = lambda c: PatternFill("solid", start_color=c, end_color=c)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = [("INFO BUNDLE", NAVY, ["Kode SKU Bundle", "Nama Produk Bundle", "Harga Bundle (Rp)"]),
        ("KOMPONEN", TEAL, ["Kode SKU Komponen", "Nama Produk Komponen", "Harga Dasar (Rp)", "Qty",
                            "Total Harga (Rp)", "Catatan Komponen"]),
        ("KEMUNGKINAN KETIDAKSESUAIAN", ORANGE,
         ["Total Harga Komponen (Rp)", "Selisih (Rp)", "Selisih (%)", "Status",
          "(1) Beli A Gratis B", "(2) Selisih Terlalu Besar", "(3) Kemungkinan Salah Harga", "Keterangan"])]
WIDTH = [16, 44, 15, 16, 44, 14, 6, 15, 24, 16, 14, 9, 22, 12, 14, 16, 70]


def analyse(master, prices, threshold, tol):
    chk, comp = build_check(master, prices, threshold)
    _, gift_c = bundle.price_of(prices, comp["component_sku"], comp["component_name"])
    comp["free_gift"] = gift_c
    chk = chk.set_index("bundle_sku")
    out = {}
    for b, c in comp.groupby("bundle_sku", sort=False):
        r = chk.loc[b]
        bp, total = r["bundle_price"], r["components_sum"]
        flags = {"a": [], "c": []}
        gap = r["status"] == "price_gap"
        if gap:
            for _, x in c.iterrows():
                u = x["unit_price"]
                if pd.isna(u) or u <= 0:
                    continue
                if x["qty"] == 1 and abs(bp - u) <= tol * u:                                  # (1) price of one component
                    flags["a"].append(x["component_sku"])
                one_unit = total - u * (x["qty"] - 1)                                          # (3) this component counted once
                if x["qty"] > 1 and bp > 0 and abs(bp - one_unit) <= tol * bp:
                    flags["c"].append(x["component_sku"])
        cond1 = bool(flags["a"]) and len(c) >= 2
        cond3 = bool(flags["c"])
        notes = []
        if cond1:
            notes.append("Harga bundle (%s) ≈ harga 1 unit %s (qty 1); komponen lain kemungkinan gratis (beli A gratis B)."
                         % (f"{bp:,.0f}", ", ".join(flags["a"])))
        if cond3:
            qtys = ", ".join(str(int(c.loc[c.component_sku == k, "qty"].iloc[0])) for k in flags["c"])
            notes.append("Harga bundle (%s) setara dengan total komponen bila %s dihitung 1 unit saja, padahal qty di "
                         "bundling = %s; kemungkinan harga bundle dihitung per unit, bukan per qty bundling." % (
                             f"{bp:,.0f}", ", ".join(flags["c"]), qtys))
        if gap and not notes:
            notes.append("Selisih harga bundle dengan total komponen %.0f%% (di atas batas %.0f%%)."
                         % (r["diff_pct"] * 100, threshold * 100))
        if r["status"] == "pass":
            notes.append("Bundle hadiah atau konversi PCS↔BOX, tidak dibandingkan.")
        if r["status"] == "bundle_price_missing":
            notes.append("Harga bundle tidak ada di file bundling price.")
        if r["status"] == "component_price_missing":
            notes.append("Ada komponen tanpa harga di file bundling price.")
        if r["status"] == "price_zero":
            notes.append("Salah satu harga (bundle atau total komponen) bernilai 0.")
        out[b] = dict(bp=bp, total=total, diff=r["diff"], pct=r["diff_pct"], status=STATUS_ID[r["status"]],
                      c1=cond1, c2=gap and not cond1, c3=cond3, notes=" ".join(notes), hit=set(flags["a"]) | set(flags["c"]),
                      flag=gap)
    return chk, comp, out


def write_sheet(ws, bundles, comp, info, title, count_label):
    ws["A1"] = title
    ws["A1"].font = Font(name="Arial", size=14, bold=True)
    ws["A2"] = count_label
    ws["A2"].font = Font(name="Arial", size=10, italic=True, color="595959")
    col = 1
    for label, color, names in HEAD:               # row 4: section band, row 5: column names
        ws.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + len(names) - 1)
        c = ws.cell(4, col, label)
        c.font, c.fill = Font(name="Arial", bold=True, color="FFFFFF"), FILL(color)
        c.alignment = Alignment(horizontal="center")
        for i, n in enumerate(names):
            h = ws.cell(5, col + i, n)
            h.font, h.fill, h.border = Font(name="Arial", bold=True, color="FFFFFF"), FILL(color), BORDER
            h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        col += len(names)
    ws.row_dimensions[5].height = 32
    for i, w in enumerate(WIDTH, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    base = Font(name="Arial", size=10)
    row = 6
    for b in bundles:
        rows = comp[comp["bundle_sku"] == b]
        i, n = info[b], len(rows)
        first, last = row, row + n - 1
        left = Alignment(vertical="top", wrap_text=True)
        # section 1 + 3 are merged over the component rows
        vals1 = [b, rows["bundle_name"].iloc[0], i["bp"]]
        vals3 = [i["total"], i["diff"], i["pct"], i["status"], "Ya" if i["c1"] else "-",
                 "Ya" if i["c2"] else "-", "Ya" if i["c3"] else "-", i["notes"]]
        for k, v in enumerate(vals1):
            _merged(ws, first, last, 1 + k, v, base, left, "#,##0" if k == 2 else None)
        for k, v in enumerate(vals3):
            fmt = {0: "#,##0", 1: "#,##0;[Red]-#,##0", 2: "0.0%"}.get(k)
            cell = _merged(ws, first, last, 10 + k, None if (isinstance(v, float) and np.isnan(v)) else v, base, left, fmt)
            if k in (4, 5, 6) and v == "Ya":
                cell.fill, cell.font = FILL("FFC7CE"), Font(name="Arial", size=10, bold=True, color="9C0006")
                cell.alignment = Alignment(horizontal="center", vertical="top")
            if k == 3:
                tint = {"Sesuai": "C6EFCE", "Perlu dicek": "FFC7CE"}.get(v, "FFEB9C" if "kosong" in v or "nol" in v else "EDEDED")
                cell.fill = FILL(tint)
        for j, (_, x) in enumerate(rows.iterrows()):
            r = first + j
            price = None if pd.isna(x["unit_price"]) else float(x["unit_price"])
            note = ("Gratis (hadiah)" if x["free_gift"] else "Harga kosong" if price is None else "")
            if "flattened" in str(x["notes"]):
                note = (note + "; " if note else "") + "Hasil penguraian bundle bersarang"
            for k, v in enumerate([x["component_sku"], x["component_name"], price, int(x["qty"]),
                                   None if price is None else price * int(x["qty"]), note]):
                c = ws.cell(r, 4 + k, v)
                c.font, c.border = base, BORDER
                c.alignment = Alignment(vertical="top", wrap_text=k in (1, 5))
                if k in (2, 4):
                    c.number_format = "#,##0"
            if price is None:
                for k in range(3, 9):
                    ws.cell(r, k + 1).fill = FILL("FFEB9C")
            if x["component_sku"] in i["hit"]:
                for k in range(3, 9):
                    ws.cell(r, k + 1).fill = FILL("FFC7CE")
        row = last + 1
    ws.freeze_panes = "D6"
    ws.sheet_view.showGridLines = False
    return row


def _merged(ws, r1, r2, c, v, font, align, fmt):
    if r2 > r1:
        ws.merge_cells(start_row=r1, start_column=c, end_row=r2, end_column=c)
    cell = ws.cell(r1, c, v)
    cell.font, cell.alignment = font, align
    if fmt:
        cell.number_format = fmt
    for r in range(r1, r2 + 1):
        ws.cell(r, c).border = BORDER
    return cell


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--bundle-dir", default=bundle.BUNDLE_DIR)
    ap.add_argument("--out", default=os.path.join(SKU_MASTER, "laporan_harga_bundling.xlsx"))
    ap.add_argument("--threshold", type=float, default=0.10, help="selisih dianggap terlalu besar (rasio)")
    ap.add_argument("--tolerance", type=float, default=0.02, help="toleransi 'sama dengan' untuk kondisi (1) dan (3)")
    a = ap.parse_args()

    prices = bundle.load_prices(a.bundle_dir)
    master = bundle.load_master(a.bundle_dir)
    chk, comp, info = analyse(master, prices, a.threshold, a.tolerance)
    order = sorted(info, key=lambda b: (not info[b]["flag"], info[b]["status"] != "Perlu dicek",
                                        -(info[b]["pct"] if not pd.isna(info[b]["pct"]) else 0), b))
    flagged = [b for b in order if info[b]["flag"]]

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Ringkasan"
    cnt = pd.Series([info[b]["status"] for b in info]).value_counts()
    ws0["A1"] = "Laporan Harga Bundling"
    ws0["A1"].font = Font(name="Arial", size=16, bold=True)
    ws0["A2"] = "Perbandingan harga bundle dengan total harga komponen (harga dasar × qty)"
    ws0["A2"].font = Font(name="Arial", size=10, italic=True, color="595959")
    lines = [("Jumlah bundle diperiksa", len(info))] + [(k, int(v)) for k, v in cnt.items()] + [
        ("", None), ("Kemungkinan ketidaksesuaian", None),
        ("(1) Beli A gratis B", sum(i["c1"] for i in info.values())),
        ("(2) Selisih terlalu besar (> %.0f%%)" % (a.threshold * 100), sum(i["c2"] for i in info.values())),
        ("(3) Kemungkinan salah harga bundle", sum(i["c3"] for i in info.values()))]
    for r, (k, v) in enumerate(lines, 4):
        ws0.cell(r, 1, k).font = Font(name="Arial", bold=(v is None and k != ""))
        if v is not None:
            ws0.cell(r, 2, v).font = Font(name="Arial")
    ws0.column_dimensions["A"].width = 46
    ws0.column_dimensions["B"].width = 12
    ws0.sheet_view.showGridLines = False

    write_sheet(wb.create_sheet("Perlu Dicek"), flagged, comp, info, "Bundle dengan Kemungkinan Harga Tidak Sesuai",
                "%d bundle" % len(flagged))
    write_sheet(wb.create_sheet("Semua Bundle"), order, comp, info, "Semua Bundle", "%d bundle" % len(order))

    ws = wb.create_sheet("Keterangan")
    text = [
        ("Bagian laporan", True),
        ("1. INFO BUNDLE: kode SKU, nama produk, dan harga bundle dari file bundling price.", False),
        ("2. KOMPONEN: satu baris per komponen (hasil penguraian bundle bersarang): harga dasar, qty, total = harga dasar × qty.", False),
        ("3. KEMUNGKINAN KETIDAKSESUAIAN: total komponen, selisih, status, dan tiga kondisi di bawah.", False),
        ("", False),
        ("Kondisi", True),
        ("(1) Beli A gratis B: harga bundle ≈ harga 1 unit salah satu komponen (qty 1); komponen lain kemungkinan gratis. Dicek hanya bila selisih melebihi batas.", False),
        ("(2) Selisih terlalu besar: |harga bundle − total komponen| ÷ harga yang lebih kecil > %.0f%%. Tidak ditandai bila bundle sudah masuk kondisi (1)." % (a.threshold * 100), False),
        ("(3) Kemungkinan salah harga: harga bundle ≈ total komponen bila satu komponen ber-qty > 1 dihitung 1 unit saja (contoh 37/390109905M = 1 unit 37/020100501A, padahal qty 6).", False),
        ("Toleransi 'sama dengan' untuk (1) dan (3): %.0f%%." % (a.tolerance * 100), False),
        ("", False),
        ("Catatan", True),
        ("Komponen kode 37/B… / nama berawalan 'B - ' (gimmick, sampling, box) dianggap hadiah gratis dengan harga 0.", False),
        ("Baris komponen berwarna merah muda = komponen yang menyerupai harga bundle; kuning = harga kosong.", False),
        ("Bundle satu komponen dengan selisih (konversi PCS↔BOX) dan bundle hadiah ditandai 'Lolos (dikecualikan)'.", False),
    ]
    for r, (t, bold) in enumerate(text, 1):
        ws.cell(r, 1, t).font = Font(name="Arial", bold=bold, size=11 if bold else 10)
    ws.column_dimensions["A"].width = 140
    ws.sheet_view.showGridLines = False

    wb.save(a.out)
    print("ditulis", a.out, "| bundle:", len(info), "| perlu dicek:", len(flagged))
    print(cnt.to_dict(), {"c1": sum(i["c1"] for i in info.values()), "c2": sum(i["c2"] for i in info.values()),
                          "c3": sum(i["c3"] for i in info.values())})


if __name__ == "__main__":
    main()
