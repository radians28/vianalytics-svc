"""Excel (Bahasa Indonesia) for confirming the bundle <-> order listing pairing with Sales.

Output : sku_master/konfirmasi_sales_bundling.xlsx
Rows   : bundle from bundling master + its components; a blank-price component ("B - ...") shows the SKU it
         refers to (a one-component bundle in the master) with that SKU's price; the order listing paired by
         product name (Tokopedia) sits on the right, followed by columns for Sales to fill in.
"""
import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from common import bundle
from common.config import SKU_MASTER, STG

# bundle -> Tokopedia platform SKU ID found by product name (None = no order listing found)
PAIRS = {"37/390709903J": "1731525346888614956", "37/390709901Q": "1731525378123858988",
         "37/390709901I": "1731525346686108716", "37/390709901A": "1731525337310135340",
         "37/390709901T": None, "37/390709903X": None}

HEAD = [("BUNDLE", "1F3864", ["Kode SKU Bundle", "Nama Bundle", "Harga Bundle (Rp)"]),
        ("KOMPONEN BUNDLE", "1F6E6E", ["Kode SKU Komponen", "Nama Komponen", "Qty", "Harga Komponen (Rp)"]),
        ("ACUAN KOMPONEN BERHARGA KOSONG", "7F6000", ["Mengacu ke SKU", "Nama SKU Acuan", "Harga SKU Acuan (Rp)"]),
        ("BERDASARKAN ORDER (TOKOPEDIA)", "9C5700",
         ["No. SKU Order (SKU ID)", "Nama Produk di Order", "Jumlah Baris Order", "Periode Order"]),
        ("KONFIRMASI SALES", "7030A0", ["Sesuai?", "Catatan Sales"])]
WIDTH = [16, 42, 14, 16, 36, 6, 14, 16, 40, 14, 22, 46, 10, 22, 14, 36]
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT = Font(name="Arial", size=10)


def main():
    m = pd.read_csv(os.path.join(bundle.BUNDLE_DIR, "bundling_master.csv"), sep=";", dtype=str,
                    encoding="utf-8-sig", keep_default_na=False)
    pr = pd.read_csv(os.path.join(bundle.BUNDLE_DIR, "bundling_price.csv"), sep=";", dtype=str, encoding="utf-8-sig",
                     keep_default_na=False).set_index("SKU")["Price"]       # as registered (blank stays blank)
    pr = pd.to_numeric(pr.replace("", pd.NA), errors="coerce")
    u = pd.read_csv(os.path.join(STG, "stg_unique_sku.csv"), dtype=str, keep_default_na=False)
    u = u[u["marketplace"] == "TOKOPEDIA"]
    price = lambda s: None if pd.isna(pr.get(s)) else float(pr.get(s))

    wb = Workbook()
    ws = wb.active
    ws.title = "Konfirmasi Bundling"
    ws["A1"] = "Konfirmasi Bundling dan Listing Order untuk Sales"
    ws["A1"].font = Font(name="Arial", size=14, bold=True)
    ws["A2"] = ("Mohon cek: apakah listing order di kanan memang bundle di kiri? Komponen berharga kosong (awalan 'B - ') "
                "mengacu ke SKU lain yang berharga; isi kolom Sesuai? dan Catatan Sales.")
    ws["A2"].font = Font(name="Arial", size=10, italic=True, color="595959")
    col = 1
    for label, color, names in HEAD:
        ws.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + len(names) - 1)
        c = ws.cell(4, col, label)
        c.font, c.fill, c.alignment = Font(name="Arial", bold=True, color="FFFFFF"), PatternFill("solid", start_color=color), Alignment(horizontal="center")
        for i, n in enumerate(names):
            h = ws.cell(5, col + i, n)
            h.font, h.border = Font(name="Arial", bold=True, color="FFFFFF"), BORDER
            h.fill = PatternFill("solid", start_color=color)
            h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        col += len(names)
    ws.row_dimensions[5].height = 32
    for i, w in enumerate(WIDTH, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    dv = DataValidation(type="list", formula1='"Sesuai,Tidak sesuai,Perlu cek"', allow_blank=True)
    ws.add_data_validation(dv)
    wrap = Alignment(vertical="top", wrap_text=True)
    row = 6
    for b, ms in PAIRS.items():
        rows = m[m["Kode SKU"] == b]
        n = len(rows)
        first, last = row, row + n - 1
        if ms:
            x = u[u["marketplace_sku"] == ms]
            order = [ms, x["product_name"].iloc[-1], int(x["order_line_count"].astype(int).sum()),
                     "%s s/d %s" % (x["first_order_at"].min()[:10], x["last_order_at"].max()[:10])]
        else:
            order = ["-", "Tidak ditemukan listing order yang sesuai", None, "-"]
        for k, v in enumerate([b, rows["Nama SKU"].iloc[0], price(b)]):
            _merge(ws, first, last, 1 + k, v, wrap, "#,##0" if k == 2 else None)
        for k, v in enumerate(order):
            _merge(ws, first, last, 11 + k, v, wrap, None)
        for k in (0, 1):
            _merge(ws, first, last, 15 + k, None, wrap, None)
        dv.add("O%d" % first)
        for j, (_, x) in enumerate(rows.iterrows()):
            r = first + j
            comp, p = x["Kode SKU BASE"], price(x["Kode SKU BASE"])
            ref = m[(m["Kode SKU"] == comp) & (m["Kode SKU BASE"] != comp)] if p is None else m.iloc[0:0]
            vals = [comp, x["Nama SKU BASE"], int(x["Qty BASE in 1 BUNDLE"] or 1), p]
            if len(ref):
                rc = ref.iloc[0]["Kode SKU BASE"]
                vals += [rc, ref.iloc[0]["Nama SKU BASE"], price(rc)]
            else:
                vals += [None, None, None]
            for k, v in enumerate(vals):
                c = ws.cell(r, 4 + k, v)
                c.font, c.border, c.alignment = FONT, BORDER, Alignment(vertical="top", wrap_text=k in (1, 5))
                if k in (3, 6):
                    c.number_format = "#,##0"
            if len(ref):
                for k in range(3, 6):
                    ws.cell(r, 4 + k + 1).fill = PatternFill("solid", start_color="FFF2CC")
        row = last + 1
    ws.freeze_panes = "D6"
    ws.sheet_view.showGridLines = False

    k = wb.create_sheet("Keterangan")
    lines = [("Tujuan", True),
             ("Memastikan listing order di Tokopedia (Feb 2023) sama dengan bundle di bundling master, sebelum dihubungkan.", False),
             ("Pencocokan listing dengan bundle berdasarkan nama produk saja (belum pasti).", False), ("", False),
             ("Komponen berharga kosong", True),
             ("Komponen berawalan 'B - ' (hadiah) tidak punya harga di file bundling price. Di bundling master SKU tersebut ternyata", False),
             ("merupakan bundle satu komponen yang mengacu ke SKU lain yang berharga (kolom 'Mengacu ke SKU').", False),
             ("Harga bundle sama persis dengan harga SKU utama (tas 2kg / 1,5kg), jadi item gratis dihitung 0 saat pembagian harga.", False),
             ("", False), ("Pertanyaan untuk Sales", True),
             ("1. Apakah listing order (kolom kanan) memang paket 'beli SKU utama gratis 2/3 pouch' sesuai bundle di kiri?", False),
             ("2. Apakah item gratis memang tidak ditagihkan (harga 0)?", False),
             ("3. Untuk 37/390709901T dan 37/390709903X belum ditemukan listing order: apakah pernah dijual, dan dengan nama apa?", False)]
    for r, (t, bold) in enumerate(lines, 1):
        k.cell(r, 1, t).font = Font(name="Arial", bold=bold, size=11 if bold else 10)
    k.column_dimensions["A"].width = 130
    k.sheet_view.showGridLines = False
    out = os.path.join(SKU_MASTER, "konfirmasi_sales_bundling.xlsx")
    wb.save(out)
    print("ditulis", out)


def _merge(ws, r1, r2, c, v, align, fmt):
    if r2 > r1:
        ws.merge_cells(start_row=r1, start_column=c, end_row=r2, end_column=c)
    cell = ws.cell(r1, c, v)
    cell.font, cell.alignment = FONT, align
    if fmt:
        cell.number_format = fmt
    for r in range(r1, r2 + 1):
        ws.cell(r, c).border = BORDER


if __name__ == "__main__":
    main()
